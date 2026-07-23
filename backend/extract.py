"""
Extract existing signal timings from the per-intersection tabs of the NYCDOT
corridor comparison workbook.

The tabs are all clones of Prototype_Tab, but they are NOT all the same width:
Kings_Hwy has 12 load-switch channels and 5 phase groups where the others have
8 and 4, which shifts every block to the right and pushes the offset row down.
So nothing here hardcodes a column letter for the timing blocks -- the layout is
discovered per sheet from these anchors:

    row 2      the literal 'Existing' / 'Proposed'  -> start of each timing block
    row 9      channel numbers 1..N (twice: existing, then proposed)
    rows 4/5   'Veh Assign' / 'Ped Assign' class per channel
    col A/B    split number and role, in 11-row bands (10 splits + 1 subtotal)
    row 6/7    plan numbers and time-of-day description
    row 16     cycle length per plan
    <offset>   the row labelled 'OFFSET'; values sit on the row below it

Only the Existing block is read. Proposed timings are out of scope.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

# Workbook layout constants that ARE stable across every tab.
ROW_BLOCK_HEADER = 2      # 'Existing' / 'Proposed'
ROW_VEH_ASSIGN = 4
ROW_PED_ASSIGN = 5
ROW_PLAN_NUMBER = 6
ROW_TOD = 7
ROW_CHANNEL_HEADER = 9
ROW_CYCLE_LENGTH = 16
ROW_FIRST_SPLIT = 18
GROUP_STRIDE = 11         # 10 split rows + 1 subtotal row per phase group
SPLITS_PER_GROUP = 10
MAX_PHASE_GROUPS = 8
MAX_SCAN_COL = 200

PHASE_LABEL_RE = re.compile(r"^PHASE\s+([A-Z])$", re.IGNORECASE)


class LayoutError(Exception):
    """The sheet does not match the expected Prototype_Tab layout."""


@dataclass
class Channel:
    channel_number: int
    kind: str                 # 'vehicle' | 'pedestrian'
    movement_class: str | None  # 'Major' | 'Minor' | None


@dataclass
class Split:
    split_number: int
    group_index: int
    position_in_group: int
    role: str | None
    indications: dict[int, str] = field(default_factory=dict)  # channel_number -> code


@dataclass
class PhaseGroup:
    group_index: int
    label: str


@dataclass
class TimingPlan:
    plan_number: int
    cycle_length_s: int
    offset_s: int
    tod_description: str | None
    durations: dict[int, int] = field(default_factory=dict)  # split_number -> seconds


@dataclass
class Intersection:
    tab_name: str
    name: str
    natural_order: int
    major_crosswalk_ft: float | None
    minor_crosswalk_ft: float | None
    channels: list[Channel]
    phase_groups: list[PhaseGroup]
    splits: list[Split]
    timing_plans: list[TimingPlan]
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------
# small cell helpers
# --------------------------------------------------------------------------

def _cell(ws, row, col):
    v = ws.cell(row=row, column=col).value
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        # Excel error literals leak through as strings when read with data_only.
        if v == "" or v.startswith("#"):
            return None
    return v


def _int(ws, row, col):
    v = _cell(ws, row, col)
    if v is None:
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def _text(ws, row, col):
    v = _cell(ws, row, col)
    return None if v is None else str(v)


# --------------------------------------------------------------------------
# layout discovery
# --------------------------------------------------------------------------

def _find_existing_block(ws) -> int:
    """Column index where the Existing timing block starts (row 2 == 'Existing')."""
    for col in range(1, MAX_SCAN_COL):
        if _text(ws, ROW_BLOCK_HEADER, col) == "Existing":
            return col
    raise LayoutError("no 'Existing' marker found in row 2")


def _find_channels(ws) -> list[Channel]:
    """
    Row 9 holds channel numbers for the Existing block, then restarts at 1 for
    the Proposed block. Read the first ascending run only.
    """
    channels: list[Channel] = []
    col = 3  # column C
    expected = 1
    while col < MAX_SCAN_COL:
        n = _int(ws, ROW_CHANNEL_HEADER, col)
        if n != expected:
            break
        veh = _text(ws, ROW_VEH_ASSIGN, col)
        ped = _text(ws, ROW_PED_ASSIGN, col)
        if veh and ped:
            raise LayoutError(
                f"channel {n} (col {get_column_letter(col)}) is assigned as both "
                f"vehicle and pedestrian"
            )
        kind = "pedestrian" if ped else "vehicle"
        channels.append(Channel(channel_number=n, kind=kind, movement_class=veh or ped))
        col += 1
        expected += 1
    if not channels:
        raise LayoutError("no channel numbers found in row 9")
    return channels


def _find_offset_label_row(ws, existing_col: int) -> int:
    """
    Locate the 'OFFSET' label, which terminates the phase-group bands.

    This must be found by scanning, not computed from the number of populated
    phase groups: every tab reserves a fixed number of 11-row bands regardless
    of how many it actually uses (Flatbush reserves 4 bands but only fills
    Phase A and Phase B), so counting used groups lands short of the real row.
    """
    for row in range(ROW_FIRST_SPLIT, ROW_FIRST_SPLIT + GROUP_STRIDE * MAX_PHASE_GROUPS + 2):
        if _text(ws, row, existing_col) == "OFFSET":
            return row
    raise LayoutError("no 'OFFSET' label found below the split rows")


def _find_phase_groups(ws, off_label_row: int) -> list[PhaseGroup]:
    """
    Each reserved band starts at ROW_FIRST_SPLIT + 11*(g-1) and is headed by a
    'PHASE X' role in column B. Unused bands are blank and are skipped, not
    treated as the end of the list.
    """
    n_bands = (off_label_row - ROW_FIRST_SPLIT) // GROUP_STRIDE
    if n_bands < 1:
        raise LayoutError(f"implausible phase-group band count ({n_bands})")

    groups: list[PhaseGroup] = []
    for g in range(n_bands):
        row = ROW_FIRST_SPLIT + GROUP_STRIDE * g
        label = _text(ws, row, 2)
        if not label:
            continue
        m = PHASE_LABEL_RE.match(label)
        if not m:
            continue
        groups.append(PhaseGroup(group_index=g + 1, label=f"Phase {m.group(1).upper()}"))
    if not groups:
        raise LayoutError("no 'PHASE X' band headers found")
    return groups


# --------------------------------------------------------------------------
# block readers
# --------------------------------------------------------------------------

def _read_splits(ws, groups: list[PhaseGroup], channels: list[Channel]) -> list[Split]:
    splits: list[Split] = []
    for grp in groups:
        base = ROW_FIRST_SPLIT + GROUP_STRIDE * (grp.group_index - 1)
        for pos in range(SPLITS_PER_GROUP):
            row = base + pos
            role = _text(ws, row, 2)
            if not role:
                break  # unused slots in the 10-row band; they are always trailing
            num = _int(ws, row, 1)
            if num is None:
                raise LayoutError(f"row {row}: role {role!r} present but no split number")
            sp = Split(
                split_number=num,
                group_index=grp.group_index,
                position_in_group=pos + 1,
                role=role,
            )
            for ch in channels:
                code = _text(ws, row, 2 + ch.channel_number)  # col C == channel 1
                if code:
                    sp.indications[ch.channel_number] = code
            splits.append(sp)
    return splits


def _read_plans(ws, existing_col: int, off_label_row: int,
                splits: list[Split]) -> list[TimingPlan]:
    off_row = off_label_row + 1

    plans: list[TimingPlan] = []
    col = existing_col
    expected = 1
    while col < MAX_SCAN_COL:
        n = _int(ws, ROW_PLAN_NUMBER, col)
        if n != expected:
            break
        cycle = _int(ws, ROW_CYCLE_LENGTH, col)
        offset = _int(ws, off_row, col)
        if not cycle:
            break  # unused plan column
        plan = TimingPlan(
            plan_number=n,
            cycle_length_s=cycle,
            offset_s=offset or 0,
            tod_description=_text(ws, ROW_TOD, col),
        )
        for sp in splits:
            base = ROW_FIRST_SPLIT + GROUP_STRIDE * (sp.group_index - 1)
            row = base + sp.position_in_group - 1
            dur = _int(ws, row, col)
            plan.durations[sp.split_number] = dur or 0
        plans.append(plan)
        col += 1
        expected += 1
    if not plans:
        raise LayoutError("no timing plans found")
    return plans


# --------------------------------------------------------------------------
# validation
# --------------------------------------------------------------------------

def _validate(inter: Intersection) -> None:
    """
    Non-fatal consistency checks, recorded on the intersection.

    The workbook enforces 'total split time equals cycle length' in the source
    controller report, so a mismatch means either a real data problem or a
    layout misread -- both worth surfacing rather than silently loading.
    """
    for plan in inter.timing_plans:
        total = sum(plan.durations.values())
        if total != plan.cycle_length_s:
            inter.warnings.append(
                f"plan {plan.plan_number}: splits sum to {total}s but cycle "
                f"length is {plan.cycle_length_s}s"
            )
        if plan.offset_s >= plan.cycle_length_s:
            inter.warnings.append(
                f"plan {plan.plan_number}: offset {plan.offset_s}s >= cycle "
                f"length {plan.cycle_length_s}s"
            )
    nums = [s.split_number for s in inter.splits]
    if nums != sorted(nums) or len(set(nums)) != len(nums):
        inter.warnings.append("split numbers are not strictly increasing/unique")


# --------------------------------------------------------------------------
# public API
# --------------------------------------------------------------------------

def extract_intersection(ws, tab_name: str, display_name: str,
                         natural_order: int) -> Intersection:
    existing_col = _find_existing_block(ws)
    channels = _find_channels(ws)
    off_label_row = _find_offset_label_row(ws, existing_col)
    groups = _find_phase_groups(ws, off_label_row)
    splits = _read_splits(ws, groups, channels)
    plans = _read_plans(ws, existing_col, off_label_row, splits)

    # Crosswalk widths live just below the timing block, labelled in column B.
    major = minor = None
    for row in range(off_label_row, off_label_row + 25):
        label = _text(ws, row, 2)
        if label == "Major" and _int(ws, row, 3) is not None:
            major = _int(ws, row, 3)
        elif label == "Minor" and _int(ws, row, 3) is not None:
            minor = _int(ws, row, 3)

    inter = Intersection(
        tab_name=tab_name,
        name=display_name,
        natural_order=natural_order,
        major_crosswalk_ft=major,
        minor_crosswalk_ft=minor,
        channels=channels,
        phase_groups=groups,
        splits=splits,
        timing_plans=plans,
    )
    _validate(inter)
    return inter


def read_intersection_index(wb) -> list[tuple[str, str, int]]:
    """
    Pair each intersection tab with its display name and spatial order.

    'IntersectionList' gives id -> display name in corridor order; the
    'IQL Tab Setup' data bank (cols Q/R/S = natural order, id, tab name) gives
    id -> tab name. Joining on id avoids guessing the mapping from tab names,
    which do not match the display names exactly (Bedford_Ave vs
    'Bedford Ave & Caton Ave').
    """
    names: dict[int, str] = {}
    ws = wb["IntersectionList"]
    for row in range(2, ws.max_row + 1):
        ident = _int(ws, row, 1)
        name = _text(ws, row, 2)
        if ident is not None and name:
            names[ident] = name

    setup = wb["IQL Tab Setup"]
    out: list[tuple[str, str, int]] = []
    seen: set[str] = set()
    for row in range(14, 400):
        order = _int(setup, row, 17)   # col Q: natural order
        ident = _int(setup, row, 18)   # col R: id
        tab = _text(setup, row, 19)    # col S: tab name
        if not tab or order is None:
            continue
        if tab in seen or tab not in wb.sheetnames:
            continue
        seen.add(tab)
        display = names.get(ident) or tab.replace("_", " ")
        out.append((tab, display, order))

    out.sort(key=lambda r: r[2])
    return out
